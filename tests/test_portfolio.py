"""Cartera: importación del historial, redes, métricas y Excel."""

import pytest

from helpers import web_client
from veritas.claims import analysis as claims
from veritas.claims import importer, service
from veritas.core.case import Case
from veritas.core.ledger import verify_chain
from veritas.forensics import network


def test_import_history_detects_rings_and_measures(history):
    work, csv_path, rep = history
    assert len(rep["imported"]) == 147 and not rep["skipped"] and not rep["ignored"]
    assert (work / "importaciones").exists() and rep["sha256"] == __import__("veritas.core.case", fromlist=["x"]).sha256_file(csv_path)
    comps = network.components(claims.load_registry(work / "registro.jsonl"))
    assert sorted(len(c["members"]) for c in comps) == [2, 4, 6]
    m = service.metrics(work)
    assert (m["tp"], m["fn"], m["fp"]) == (9, 3, 4)          # detección 75%, precisión 69%
    assert m["recall"] == 75 and m["precision"] == 69
    case = Case(work / rep["imported"][0][0])
    actions = [e["action"] for e in case.ledger.entries()]
    assert "imported_from" in actions and verify_chain(case.ledger.entries())[0]


def test_import_reimport_skips_existing_and_bad_rows(history, tmp_path):
    work, csv_path, _ = history
    rep = importer.import_history(work, csv_path)
    assert not rep["imported"] and len(rep["skipped"]) == 147
    bad = tmp_path / "malo.csv"
    bad.write_text("Siniestro;Fecha;Monto\nX-1;31-02-2026;100\n;20-09-2026;5\nX-3;20/09/2026 10:15;$1.000\n", encoding="utf-8")
    rep = importer.import_history(tmp_path / "w", bad)
    assert [n for n, _, _ in rep["imported"]] == ["X-3"] and len(rep["skipped"]) == 2


@pytest.mark.filterwarnings("ignore::EncodingWarning")  # interna de openpyxl al guardar el archivo de prueba
def test_import_excel_with_windows_headers(tmp_path):
    from openpyxl import Workbook
    from datetime import datetime as dt
    wb = Workbook()
    ws = wb.active
    ws.append(["Nº Siniestro", "Fecha del siniestro", "RUT", "Cuenta de pago", "Relato", "Estado final"])
    ws.append(["E-1", dt(2026, 5, 1, 12, 0), "1-9", "123", "Choque leve en estacionamiento de supermercado sin testigos", "Pagado"])
    ws.append(["E-2", dt(2026, 6, 1, 12, 0), "2-7", "123", "Otro choque en la calle principal del barrio", "Fraude"])
    path = tmp_path / "h.xlsx"
    wb.save(path)
    rep = importer.import_history(tmp_path / "w", path)
    assert len(rep["imported"]) == 2
    assert service.current_decision(Case(tmp_path / "w" / "E-2"))["subject"] == "fraude"
    assert service.last_analysis(Case(tmp_path / "w" / "E-2"))["level"] == "bad"  # misma cuenta, otro RUT


def test_web_import_networks_metrics_decision(tmp_path):
    import io as _io
    c = web_client(tmp_path)
    csv_bytes = ("N° Siniestro;Fecha;RUT;Teléfono;Resultado\nW-1;01-05-2026 10:00;1-9;+56 9 1111 2222;fraude\n"
                 "W-2;02-06-2026 11:00;2-7;+56 9 1111 2222;pagado\nW-3;03-07-2026 12:00;3-5;+56 9 1111 2222;\n").encode("cp1252")
    r = c.post("/importar", data={"actor": "Prueba", "archivo": (_io.BytesIO(csv_bytes), "hist.csv")}, content_type="multipart/form-data")
    page = r.get_data(as_text=True)
    assert r.status_code == 200 and "3</b><span>siniestros importados" in page
    redes = c.get("/redes").get_data(as_text=True)
    assert "Red de 3 siniestros" in redes and "mismo teléfono" in redes
    met = c.get("/metricas").get_data(as_text=True)
    assert "Aciertos y errores" in met
    r = c.post("/caso/W-3/decision", data={"decision": "fraude", "actor": "Ana", "nota": "ok"}, follow_redirects=True)
    assert "Decisión registrada: Fraude confirmado" in r.get_data(as_text=True)
    assert "Fraude confirmado" in c.get("/siniestros").get_data(as_text=True)
    assert c.get("/plantilla.csv").get_data(as_text=True).lstrip("﻿").startswith("numero;fecha_siniestro")
    r = c.post("/importar", data={"archivo": (_io.BytesIO(b"x"), "hist.txt")}, content_type="multipart/form-data", follow_redirects=True)
    assert "Seleccione un archivo CSV o Excel" in r.get_data(as_text=True)


def test_excel_exports(loaded):
    import io as _io
    from openpyxl import load_workbook
    from veritas.web import create_app
    c = create_app(loaded[0]).test_client()
    for what, sheet in (("cola", "Cola de trabajo"), ("redes", "Redes"), ("metricas", "Resumen")):
        r = c.get(f"/exportar/{what}.xlsx")
        assert r.status_code == 200 and sheet in load_workbook(_io.BytesIO(r.data)).sheetnames
    assert c.get("/exportar/otra.xlsx").status_code == 404
