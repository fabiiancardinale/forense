"""Seguridad de los módulos de siniestros: revisión de archivos, análisis aislado, límite del portal y Excel."""
import io
import zipfile

import pytest
from PIL import Image

from helpers import LIMPIO, real_pdf, textured_image, web_client
from evidex.claims import isolated, service
from evidex.core import upload_guard
from evidex.inspection.validation import InvalidFile


def _jpeg() -> bytes:
    b = io.BytesIO()
    textured_image(3, 400, 300).save(b, "JPEG", quality=90)
    return b.getvalue()


def test_guard_checks_real_content(tmp_path):
    def f(name, data):
        p = tmp_path / name
        p.write_bytes(data)
        return p
    assert upload_guard.check(f("a.jpg", _jpeg()), "a.jpg", "evidencia")["format"] == "JPEG"
    assert upload_guard.check(f("d.pdf", real_pdf()), "d.pdf", "informe")["kind"] == "pdf"
    b = io.BytesIO()
    Image.new("RGB", (40, 40)).save(b, "TIFF")
    assert upload_guard.check(f("s.tif", b.getvalue()), "s.tif", "documento")["format"] == "TIFF"
    rejects = [
        ("falsa.jpg", b"<script>alert(1)</script>", "evidencia"),       # no es una foto
        ("falso.pdf", b"%PDF-1.4\nbasura", "portal"),                   # no es un PDF sano
        ("a.jpg", _jpeg(), "informe"),                                  # el informe del perito es solo PDF
        ("nota.txt", b"hola", "portal"),                                # el asegurado no sube textos
        ("vacio.png", b"", "evidencia"),
    ]
    for name, data, purpose in rejects:
        with pytest.raises(InvalidFile):
            upload_guard.check(f(name, data), name, purpose)
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zz:
        zz.writestr("../../fuera.txt", "x")
    with pytest.raises(InvalidFile, match="rutas"):
        upload_guard.check(f("chat.zip", z.getvalue()), "chat.zip", "evidencia")
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zz:
        zz.writestr("a.txt", b"0" * 30_000_000)
    with pytest.raises(InvalidFile, match="compresión"):
        upload_guard.check(f("bomba.zip", z.getvalue()), "bomba.zip", "evidencia")


def test_web_rejects_fake_files_with_reason(tmp_path):
    c = web_client(tmp_path)
    r = c.post("/nuevo", data={"numero": "X1", "fecha_siniestro": "2026-09-20T18:30",
                                "fotos": [(io.BytesIO(b"<html>no soy foto</html>"), "danio.jpg")]},
               content_type="multipart/form-data")
    page = r.get_data(as_text=True)
    assert r.status_code == 400 and "danio.jpg" in page and "Imagen inválida" in page
    r = c.post("/auditar", data={"pdf": (io.BytesIO(b"%PDF-1.4 falso"), "informe.pdf")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "Informe rechazado" in r.get_data(as_text=True)
    r = c.post("/importar", data={"archivo": (io.BytesIO(b"PK\x03\x04roto"), "h.xlsx")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "Archivo rechazado" in r.get_data(as_text=True)


def test_portal_rejects_fake_file_and_limits_guessing(tmp_path):
    from evidex.portal import links as capture
    from evidex.web import create_app
    service.load_demo(tmp_path)
    case = service.find_case(tmp_path, LIMPIO)
    app = create_app(tmp_path, {"DEMO_MODE": True, "CSRF_ENABLED": False, "LEGACY_MODULES": True,
                                "PORTAL_MAX_MISSES": 5})
    c = app.test_client()
    c.post(f"/caso/{case.root.name}/captura")
    tok = capture.load(case)["token"]
    r = c.post(f"/c/{tok}/archivo", content_type="multipart/form-data",
               data={"kind": "doc", "titulo": "Licencia", "archivo": (io.BytesIO(b"%PDF-1.4 falso"), "lic.pdf")})
    assert r.status_code == 400 and "no se aceptó" in r.get_json()["error"]
    r = c.post(f"/c/{tok}/archivo", content_type="multipart/form-data",
               data={"kind": "doc", "titulo": "Licencia", "archivo": (io.BytesIO(real_pdf()), "lic.pdf")})
    assert r.status_code == 200
    for i in range(5):                                    # alguien probando enlaces inventados
        assert c.get("/c/" + f"{i}" * 30).status_code == 404
    assert c.get(f"/c/{tok}").status_code == 429          # bloqueado, aunque ahora use un enlace real
    other = c.get(f"/c/{tok}", headers={"Cf-Connecting-Ip": "200.1.2.3"})
    assert other.status_code == 200                       # detrás del túnel se cuenta por la IP real


def test_isolated_analysis_and_failure_keeps_case_usable(tmp_path, monkeypatch):
    service.load_demo(tmp_path)
    case = service.find_case(tmp_path, LIMPIO)
    r = isolated.analyze(case, tmp_path / "registro.jsonl")
    assert r["findings"] >= 0 and "recommendation" in r
    with pytest.raises(isolated.AnalysisFailed, match="tiempo máximo"):
        isolated.analyze(case, tmp_path / "registro.jsonl", timeout=0.01)

    from evidex.web import create_app
    app = create_app(tmp_path, {"DEMO_MODE": True, "CSRF_ENABLED": False, "LEGACY_MODULES": True})
    c = app.test_client()

    def boom(*a, **kw):
        raise isolated.AnalysisFailed("El análisis superó el tiempo máximo y se detuvo.")
    monkeypatch.setattr(isolated, "analyze", boom)
    r = c.post(f"/caso/{case.root.name}/analizar", follow_redirects=True)
    page = r.get_data(as_text=True)
    assert r.status_code == 200 and "conserva su análisis anterior" in page


def test_excel_never_writes_formulas():
    from openpyxl import Workbook, load_workbook
    from evidex.claims.excel import _bytes, _sheet
    wb = Workbook()
    _sheet(wb, "X", ["Relato"], [['=HYPERLINK("http://malo","clic")'], ["texto normal"]], [20])
    ws = load_workbook(io.BytesIO(_bytes(wb)))["X"]
    assert [c.data_type for c in ws["A"]] == ["s", "s", "s"]


def test_full_app_is_default_from_command_line(monkeypatch):
    from evidex.web import __main__ as cli
    seen = {}

    def fake_create(workdir, config):
        seen.update(config)
        raise SystemExit(0)
    monkeypatch.setattr(cli, "create_app", fake_create)
    with pytest.raises(SystemExit):
        cli.main(["--dir", "x", "--no-abrir"])
    assert seen["LEGACY_MODULES"] is True
    with pytest.raises(SystemExit):
        cli.main(["--dir", "x", "--no-abrir", "--solo-analisis"])
    assert seen["LEGACY_MODULES"] is False


def test_temporary_address_warns_and_detects_changed_link(tmp_path):
    from evidex.portal import links as capture
    from evidex.web import create_app
    service.load_demo(tmp_path)
    case = service.find_case(tmp_path, LIMPIO)
    app = create_app(tmp_path, {"DEMO_MODE": True, "CSRF_ENABLED": False, "LEGACY_MODULES": True,
                                "PUBLIC": True, "PUBLIC_BASE": "https://uno-dos-tres.trycloudflare.com"})
    c = app.test_client()
    c.post(f"/caso/{case.root.name}/captura")
    tab = c.get(f"/caso/{case.root.name}?tab=asegurado").get_data(as_text=True)
    assert "Dirección temporal" in tab and "data-sent-url" in tab
    assert c.post(f"/caso/{case.root.name}/captura/enviado").status_code == 204
    assert capture.load(case)["sent_base"] == "https://uno-dos-tres.trycloudflare.com"
    app.config["PUBLIC_BASE"] = "https://cuatro-cinco.trycloudflare.com"      # Evidex se reinició: otro túnel
    tab = c.get(f"/caso/{case.root.name}?tab=asegurado").get_data(as_text=True)
    assert "El enlace que envió ya no funciona" in tab and "uno-dos-tres" in tab
    assert "Reenviar enlace" in c.get("/").get_data(as_text=True)
    c.post(f"/caso/{case.root.name}/captura/enviado")                         # lo reenvió
    assert "ya no funciona" not in c.get(f"/caso/{case.root.name}?tab=asegurado").get_data(as_text=True)
    assert capture.share_message("X", "alberto pérez", "https://a/c/t", "2026-10-08T10:00:00").startswith("Hola Alberto")


def test_case_opens_immediately_and_analyzes_in_background(tmp_path, monkeypatch):
    import time
    from evidex.web import background, create_app
    from evidex.claims.service import find_case
    service.load_demo(tmp_path)
    case = find_case(tmp_path, LIMPIO)
    (case.root / "informe.evidencias").write_text("0|viejo")                 # como después de actualizar Evidex
    app = create_app(tmp_path, {"DEMO_MODE": True, "CSRF_ENABLED": False, "LEGACY_MODULES": True,
                                "BACKGROUND_ANALYSIS": True})
    c = app.test_client()
    t = time.monotonic()
    page = c.get(f"/caso/{case.root.name}").get_data(as_text=True)
    assert time.monotonic() - t < 2 and "Analizando las pruebas" in page     # no espera al análisis
    assert c.get(f"/caso/{case.root.name}/analisis").get_json()["running"]
    background.wait(case, 120)
    assert c.get(f"/caso/{case.root.name}/analisis").get_json() == {
        "running": False, "state": "done", "finished": background.status(case)["finished"]}
    assert "Analizando las pruebas" not in c.get(f"/caso/{case.root.name}").get_data(as_text=True)

    # si falla, se avisa y no se reintenta solo en cada visita
    def boom(*a, **kw):
        raise isolated.AnalysisFailed("El análisis superó el tiempo máximo y se detuvo.")
    monkeypatch.setattr(isolated, "analyze", boom)
    (case.root / "informe.evidencias").write_text("0|viejo")
    c.get(f"/caso/{case.root.name}")
    background.wait(case, 30)
    page = c.get(f"/caso/{case.root.name}").get_data(as_text=True)
    assert "El último análisis no terminó" in page and not background.is_running(case)


def test_tunnel_allows_portal_design_only(tmp_path):
    from evidex.web import create_app
    app = create_app(tmp_path, {"LEGACY_MODULES": True, "PUBLIC": True, "PUBLIC_BASE": "https://x.trycloudflare.com"})
    c = app.test_client()
    via = {"Cf-Ray": "x", "Cf-Connecting-Ip": "200.1.1.1"}
    assert c.get("/static/portal.css", headers=via).status_code == 200
    for path in ("/static/app.css", "/static/inspection.js", "/ingresar", "/"):
        assert c.get(path, headers=via).status_code == 404, path
