import json
import subprocess
import sys
import zipfile
from pathlib import Path

import os

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
UTF8_ENV = {**os.environ, "PYTHONIOENCODING": "utf-8"}

from veritas import assistant, detections  # noqa: E402
from veritas.case import Case  # noqa: E402
from veritas.cli import main  # noqa: E402
from veritas.ledger import verify_chain  # noqa: E402
from veritas.timeline import Timeline  # noqa: E402


@pytest.fixture
def built(tmp_path):
    subprocess.run([sys.executable, str(ROOT / "demo/make_demo.py"), str(tmp_path / "ev.jsonl")], check=True, capture_output=True)
    case_dir = tmp_path / "caso"
    main(["init", str(case_dir), "Incidente demo"])
    main(["add", str(case_dir), str(tmp_path / "ev.jsonl")])
    main(["ingest", str(case_dir)])
    return case_dir


def test_detects_full_attack_chain(built):
    tl = Timeline(built / "timeline.db")
    rules = {f.rule for f in detections.run_all(tl.all())}
    assert {"brute_force_success", "encoded_powershell", "admin_added", "service_installed", "log_cleared"} <= rules


def test_ledger_tamper_is_detected(built):
    case = Case(built)
    assert verify_chain(case.ledger.entries())[0]
    lines = (built / "ledger.jsonl").read_text(encoding="utf-8").splitlines()
    e = json.loads(lines[1]); e["actor"] = "otro"
    lines[1] = json.dumps(e)
    (built / "ledger.jsonl").write_text("\n".join(lines) + "\n", encoding="utf-8")
    assert not verify_chain(Case(built).ledger.entries())[0]


def test_evidence_tamper_is_detected(built):
    case = Case(built)
    f = next(case.evidence_dir.iterdir())
    f.chmod(0o644); f.write_text("manipulado", encoding="utf-8")
    ok, problems = case.verify_evidence()
    assert not ok and problems


def test_assistant_rejects_ungrounded_claims(built):
    tl = Timeline(built / "timeline.db")
    real = tl.all()[0]["id"]
    text = f"Hubo un acceso [ev:{real}]. El atacante robó la base de datos. Hecho falso [ev:deadbeef:999]."
    v = assistant.verify_claims(text, tl)
    assert len(v.accepted) == 1 and len(v.rejected) == 2


def test_export_bundle_verifies_and_detects_tamper(built, tmp_path):
    z = tmp_path / "paquete.zip"
    main(["report", str(built)])
    main(["export", str(built), str(z)])
    out = tmp_path / "pkg"
    zipfile.ZipFile(z).extractall(out)
    r = subprocess.run([sys.executable, str(out / "verify.py")], capture_output=True, text=True, encoding="utf-8", env=UTF8_ENV)
    assert r.returncode == 0, r.stdout
    assert "firma digital válida" in r.stdout
    # alterar una entrada y volver a verificar
    led = out / "ledger.jsonl"
    led.write_text(led.read_text(encoding="utf-8").replace("case_created", "case_deleted"), encoding="utf-8")
    r = subprocess.run([sys.executable, str(out / "verify.py")], capture_output=True, text=True, encoding="utf-8", env=UTF8_ENV)
    assert r.returncode == 1


def test_offline_summary_keeps_claims(built):
    tl = Timeline(built / "timeline.db")
    v = assistant.summarize(detections.run_all(tl.all()), tl)
    assert len(v.accepted) >= 5 and not v.rejected
    assert all(assistant.CITE.sub("", s).strip(" .") for s in v.accepted)


# ---- siniestros -------------------------------------------------------------
from veritas import claims, documents, narrative, network, policy, service  # noqa: E402

ACTUAL = "SIN-2026-0987"
LIMPIO = "SIN-2026-0990"


@pytest.fixture(scope="module")
def demo_src(tmp_path_factory):
    d = tmp_path_factory.mktemp("demo") / "src"
    subprocess.run([sys.executable, str(ROOT / "demo/make_claim_demo.py"), str(d)], check=True, capture_output=True)
    return d


@pytest.fixture(scope="module")
def loaded(tmp_path_factory):
    work = tmp_path_factory.mktemp("work") / "casos"
    results = dict(service.load_demo(work))
    return work, results


def _analysis(work, numero):
    return service.run_analysis(Case(work / numero), work / "registro.jsonl")


def test_demo_loads_in_order_with_expected_outcome(loaded):
    _, r = loaded
    assert list(r) == ["SIN-2026-0412", "SIN-2026-0533", "SIN-2026-0601", ACTUAL, LIMPIO]
    assert r[ACTUAL]["level"] == "bad" and r[ACTUAL]["linked"] == 3 and r[ACTUAL]["score"] == 100
    assert r[LIMPIO]["findings"] == 0 and r[LIMPIO]["level"] == "ok"
    assert r["SIN-2026-0533"]["level"] == "bad"  # misma cuenta y relato copiado del 0412
    # al aparecer los siniestros nuevos, los históricos vinculados se reanalizaron
    assert r["SIN-2026-0533"]["reanalyzed"] == ["SIN-2026-0412"]
    assert set(r[ACTUAL]["reanalyzed"]) == {"SIN-2026-0412", "SIN-2026-0533", "SIN-2026-0601"}
    work = loaded[0]
    last = service.last_analysis(Case(work / "SIN-2026-0601"))
    assert last["level"] == "bad" and last["trigger"] == f"reanálisis por vínculo con {ACTUAL}"
    assert service.last_analysis(Case(work / LIMPIO))["level"] == "ok"


def test_suspicious_claim_triggers_every_module(loaded):
    work, _ = loaded
    findings = _analysis(work, ACTUAL)[4]
    rules = {f.rule for f in findings}
    expected = {
        "edited", "reused_photo", "no_metadata", "ai_generated", "taken_before", "gps_mismatch",       # fotos
        "doc_editor", "doc_modified", "doc_before_claim", "reused_doc",                                 # documentos
        "coverage_change", "early_claim", "late_report", "near_limit",                                  # póliza
        "copied_narrative", "time_contradiction", "police_missing",                                     # relato
        "shared_contact", "fraud_ring",                                                                 # red
        "doc_version_changed", "doc_rut_invalid",                                                       # versiones y RUT
        "impossible_travel", "photo_before_policy",                                                     # línea de tiempo
        "chat_before_event", "chat_script", "chat_deleted", "chat_phone_link",                          # chat
    }
    assert expected <= rules, expected - rules
    assert not any("01_trasera" in f.title for f in findings)  # la foto legítima no se marca
    assert {f.category for f in findings} == {"Fotos", "Documentos", "Póliza y siniestro", "Relato", "Red de siniestros",
                                              "Línea de tiempo", "Comunicaciones", "Datos del asegurado"}


def test_clean_claim_has_no_alerts(loaded):
    work, _ = loaded
    assert _analysis(work, LIMPIO)[4] == []
    assert "Sin alertas" in (work / LIMPIO / "informe.html").read_text(encoding="utf-8")


def test_report_has_network_and_grounded_summary(loaded):
    work, _ = loaded
    html = (work / ACTUAL / "informe.html").read_text(encoding="utf-8")
    assert '<svg class="net"' in html and "SIN-2026-0601" in html
    assert "0 afirmación(es) sin evidencia" in html
    for section in ("Red de siniestros", "Póliza y siniestro", "Documentos", "Relato del asegurado"):
        assert section in html


def test_export_bundle_of_claim_verifies(loaded, tmp_path):
    work, _ = loaded
    z = tmp_path / "s.zip"
    main(["export", str(work / ACTUAL), str(z)])
    zipfile.ZipFile(z).extractall(tmp_path / "pkg")
    r = subprocess.run([sys.executable, str(tmp_path / "pkg/verify.py")], capture_output=True, text=True, encoding="utf-8", env=UTF8_ENV)
    assert r.returncode == 0, r.stdout
    assert "presupuesto_taller.pdf" in r.stdout


def test_only_the_intended_photo_is_a_duplicate(demo_src):
    from PIL import Image
    fs = sorted(demo_src.rglob("fotos/*"))
    hs = {f.name: claims.dhash(Image.open(f)) for f in fs}
    close = {tuple(sorted((a, b))) for a in hs for b in hs if a < b and claims.hamming(hs[a], hs[b]) <= claims.DUP_MAX_DISTANCE}
    assert close == {("03_parachoques.jpg", "IMG_4471.jpg")}


# ---- módulos por separado ------------------------------------------------------
def test_policy_rules_and_amount_parsing():
    assert policy.parse_amount("$4.650.000") == 4650000 == policy.parse_amount("4650000.0") == policy.parse_amount(4650000)
    base = {"numero": "X", "fecha_siniestro": "2026-09-20T03:00:00"}
    rules = {f.rule for f in policy.analyze({**base, "inicio_poliza": "2026-09-25", "fecha_denuncia": "2026-11-30"}, "e:1")}
    assert {"before_policy", "late_report", "night_no_witness"} <= rules
    calm = {**base, "fecha_siniestro": "2026-09-20T15:00:00", "inicio_poliza": "2024-01-01", "fin_poliza": "2027-01-01",
            "fecha_denuncia": "2026-09-21", "suma_asegurada": "5.000.000", "monto_reclamado": "900.000"}
    assert policy.analyze(calm, "e:1") == []


def test_narrative_copy_and_contradictions():
    a = "Estaba detenido en el semáforo cuando un auto me chocó por detrás y se dio a la fuga sin dejar sus datos."
    b = "Me encontraba detenido en el semáforo cuando un auto me chocó por detrás y se dio a la fuga sin dejar datos."
    c = "Al salir del estacionamiento del supermercado otro vehículo retrocedió y golpeó la puerta del conductor."
    assert narrative.similarity(a, b) > 0.5 > narrative.similarity(a, c)
    decl = {"numero": "N", "fecha_siniestro": "2026-01-01T14:00:00", "testigos": 2,
            "descripcion": "Ocurrió en la noche, no hubo testigos y dejé constancia en Carabineros de lo que pasó ese día."}
    rules = {f.rule for f in narrative.analyze(decl, "e:1", {"OTRO": {"relato": decl["descripcion"]}})}
    assert {"copied_narrative", "time_contradiction", "witness_contradiction", "police_missing"} <= rules


def test_pdf_metadata_detects_incremental_edit(demo_src, tmp_path):
    m = documents.pdf_metadata(demo_src / "actual/documentos/presupuesto_taller.pdf")
    assert m["is_pdf"] and m["revisions"] == 2 and "iLovePDF" in m["tools"]
    clean = documents.pdf_metadata(demo_src / "limpio/documentos/presupuesto_automotriz_central.pdf")
    assert clean["revisions"] == 1 and clean["created"].startswith("2026-09-22")
    fake = tmp_path / "factura.pdf"
    fake.write_bytes(b"esto no es un pdf")
    d = documents.Document("ab" * 32, "factura.pdf", documents.pdf_metadata(fake), "2026-09-21T00:00:00")
    assert [f.rule for f in documents.analyze({"numero": "N", "fecha_siniestro": "2026-09-20T10:00:00"}, "e:1", [d], [])] == ["not_pdf"]


def test_network_normalization_and_same_person_is_not_a_ring():
    assert network.normalize("rut", "12.345.678-k") == "12345678K"
    assert network.normalize("telefono", "+56 9 4444 9090") == network.normalize("telefono", "944449090")
    # el mismo asegurado con varios siniestros: frecuencia sí, red de personas distintas no
    reg = [network.claim_record({"numero": f"S{i}", "rut": "11.111.111-1", "telefono": "+56 9 1111 2222",
                                 "fecha_siniestro": f"2026-0{i}-10T10:00:00"}, "ok") for i in (3, 5, 7)]
    decl = {"numero": "S9", "rut": "11111111-1", "telefono": "91111 2222", "fecha_siniestro": "2026-09-10T10:00:00"}
    f, net = network.analyze(decl, "e:1", [], [], reg)
    assert {x.rule for x in f} == {"frequency_rut"} and net["group"] == []


# ---- línea de comandos -----------------------------------------------------------
def test_cli_accepts_pdf_and_rejects_other_files(demo_src, tmp_path):
    src = demo_src / "limpio"
    case = tmp_path / "c"
    main(["siniestro-init", str(case), str(src / "declaracion.json"), str(src / "fotos/a_frontal.jpg"),
          str(src / "documentos/presupuesto_automotriz_central.pdf")])
    main(["siniestro-analizar", str(case)])
    assert "presupuesto_automotriz_central.pdf" in (case / "informe.html").read_text(encoding="utf-8")
    (tmp_path / "nota.txt").write_text("x", encoding="utf-8")
    with pytest.raises(SystemExit, match="no son fotos ni PDF"):
        main(["siniestro-init", str(tmp_path / "d"), str(src / "declaracion.json"), str(tmp_path / "nota.txt")])


def test_declaration_saved_as_windows_ansi_is_accepted(demo_src, tmp_path):
    decl = json.loads((demo_src / "limpio/declaracion.json").read_text(encoding="utf-8"))
    ansi = tmp_path / "decl_ansi.json"
    ansi.write_bytes(json.dumps(decl, ensure_ascii=False).encode("cp1252"))  # como el Bloc de notas
    case = tmp_path / "ansi"
    main(["siniestro-init", str(case), str(ansi), str(demo_src / "limpio/fotos/a_frontal.jpg")])
    main(["siniestro-analizar", str(case)])
    assert "camión de reparto" in (case / "informe.html").read_text(encoding="utf-8")


def test_clear_errors_for_missing_or_duplicate_case(demo_src, tmp_path):
    decl = str(demo_src / "limpio/declaracion.json")
    with pytest.raises(SystemExit, match="No existe el caso"):
        main(["siniestro-analizar", str(tmp_path / "no_existe")])
    main(["siniestro-init", str(tmp_path / "dup"), decl])
    with pytest.raises(SystemExit, match="ya existe"):
        main(["siniestro-init", str(tmp_path / "dup"), decl])
    with pytest.raises(SystemExit, match="No se encuentra"):
        main(["siniestro-init", str(tmp_path / "x"), str(tmp_path / "falta.json")])


# ---- interfaz web -------------------------------------------------------------
def _web(tmp_path):
    from veritas.web import create_app
    app = create_app(tmp_path / "casos")
    app.config["TESTING"] = True
    return app.test_client()


def test_web_create_claim_with_documents_verify_export(demo_src, tmp_path):
    import io as _io
    c = _web(tmp_path)
    assert "Aún no hay siniestros" in c.get("/").get_data(as_text=True)
    src = demo_src / "actual"
    decl = json.loads((src / "declaracion.json").read_text(encoding="utf-8"))
    form = {k: str(v) for k, v in decl.items() if k not in ("lat", "lon", "fecha_siniestro")}
    form.update({"numero": "SIN-2026-5000", "fecha_siniestro": "2026-09-20T18:30", "lat": "-33,4263", "lon": "-70.6167"})
    files = sorted((src / "fotos").iterdir()) + sorted((src / "documentos").iterdir())
    handles = [(open(p, "rb"), p.name) for p in files]
    r = c.post("/nuevo", data={**form, "fotos": handles}, content_type="multipart/form-data")
    for fh, _ in handles:
        fh.close()
    assert r.status_code == 302 and r.location.endswith("/caso/SIN-2026-5000")
    assert "Derivar a la unidad de investigación" in c.get("/caso/SIN-2026-5000").get_data(as_text=True)
    rep = c.get("/caso/SIN-2026-5000/informe").get_data(as_text=True)
    assert "Documento pasado por un editor de PDF" in rep and "Aumento de cobertura" in rep
    index = c.get("/siniestros").get_data(as_text=True)
    assert "SIN-2026-5000" in index and 'class="risk bad"' in index
    dash = c.get("/").get_data(as_text=True)
    assert "Derivados sin decisión" in dash and "SIN-2026-5000" in dash
    assert "Integridad verificada" in c.post("/caso/SIN-2026-5000/verificar", follow_redirects=True).get_data(as_text=True)
    r = c.get("/caso/SIN-2026-5000/exportar")
    assert r.mimetype == "application/zip" and "verify.py" in zipfile.ZipFile(_io.BytesIO(r.data)).namelist()
    with open(src / "fotos/01_trasera.jpg", "rb") as fh:
        r = c.post("/caso/SIN-2026-5000/fotos", data={"fotos": [(fh, "01_trasera.jpg")]},
                   content_type="multipart/form-data", follow_redirects=True)
    assert "ya estaban en el caso" in r.get_data(as_text=True)


def test_web_validation_and_demo(tmp_path):
    import io as _io
    c = _web(tmp_path)
    r = c.post("/nuevo", data={"numero": "X1", "fecha_siniestro": "2026-09-20T18:30",
                                "fotos": [(_io.BytesIO(b"hola"), "nota.txt")]}, content_type="multipart/form-data")
    assert r.status_code == 400 and "no son fotos ni PDF" in r.get_data(as_text=True)
    r = c.post("/nuevo", data={"numero": "X1", "fecha_siniestro": "2026-09-20T18:30"}, content_type="multipart/form-data")
    assert "al menos una foto o documento" in r.get_data(as_text=True)
    assert c.get("/caso/..%2Fetc").status_code == 404
    page = c.post("/demo", follow_redirects=True).get_data(as_text=True)
    assert all(n in page for n in ("SIN-2026-0412", "SIN-2026-0533", "SIN-2026-0601", ACTUAL, LIMPIO))
    assert "Derivar a la unidad" in page and "Sin alertas" in page
    assert "ya estaba cargado" in c.post("/demo", follow_redirects=True).get_data(as_text=True)


# ---- historial, redes globales, decisiones y métricas --------------------------------
from veritas import importer  # noqa: E402


@pytest.fixture(scope="module")
def history(tmp_path_factory):
    d = tmp_path_factory.mktemp("hist")
    csv_path = d / "historial.csv"
    subprocess.run([sys.executable, str(ROOT / "demo/make_history_demo.py"), str(csv_path)], check=True, capture_output=True)
    work = d / "casos"
    return work, csv_path, importer.import_history(work, csv_path, actor="prueba")


def test_import_history_detects_rings_and_measures(history):
    work, csv_path, rep = history
    assert len(rep["imported"]) == 147 and not rep["skipped"] and not rep["ignored"]
    assert (work / "importaciones").exists() and rep["sha256"] == __import__("veritas.case", fromlist=["x"]).sha256_file(csv_path)
    comps = network.components(claims.load_registry(work / "registro.jsonl"))
    assert sorted(len(c["members"]) for c in comps) == [2, 4, 6]
    m = service.metrics(work)
    assert (m["tp"], m["fn"], m["fp"]) == (9, 3, 4)          # detección 75%, precisión 69%
    assert m["recall"] == 75 and m["precision"] == 69
    case = Case(work / rep["imported"][0][0])
    actions = [e["action"] for e in case.ledger.entries()]
    assert "imported_from" in actions and verify_chain(case.ledger.entries())[0]


def test_import_reimport_skips_existing_and_bad_rows(history, tmp_path):
    work, csv_path, _ = history
    rep = importer.import_history(work, csv_path)
    assert not rep["imported"] and len(rep["skipped"]) == 147
    bad = tmp_path / "malo.csv"
    bad.write_text("Siniestro;Fecha;Monto\nX-1;31-02-2026;100\n;20-09-2026;5\nX-3;20/09/2026 10:15;$1.000\n", encoding="utf-8")
    rep = importer.import_history(tmp_path / "w", bad)
    assert [n for n, _, _ in rep["imported"]] == ["X-3"] and len(rep["skipped"]) == 2


@pytest.mark.filterwarnings("ignore::EncodingWarning")  # interna de openpyxl al guardar el archivo de prueba
def test_import_excel_with_windows_headers(tmp_path):
    from openpyxl import Workbook
    from datetime import datetime as dt
    wb = Workbook()
    ws = wb.active
    ws.append(["Nº Siniestro", "Fecha del siniestro", "RUT", "Cuenta de pago", "Relato", "Estado final"])
    ws.append(["E-1", dt(2026, 5, 1, 12, 0), "1-9", "123", "Choque leve en estacionamiento de supermercado sin testigos", "Pagado"])
    ws.append(["E-2", dt(2026, 6, 1, 12, 0), "2-7", "123", "Otro choque en la calle principal del barrio", "Fraude"])
    path = tmp_path / "h.xlsx"
    wb.save(path)
    rep = importer.import_history(tmp_path / "w", path)
    assert len(rep["imported"]) == 2
    assert service.current_decision(Case(tmp_path / "w" / "E-2"))["subject"] == "fraude"
    assert service.last_analysis(Case(tmp_path / "w" / "E-2"))["level"] == "bad"  # misma cuenta, otro RUT


def test_decisions_are_in_custody_and_feed_metrics(tmp_path, demo_src):
    work = tmp_path / "w"
    service.load_demo(work)
    case = Case(work / ACTUAL)
    service.set_decision(case, "fraude", "Liquidador Prueba", "Confirmado con el taller")
    service.set_decision(Case(work / LIMPIO), "legitimo", "Liquidador Prueba")
    assert service.current_decision(case)["actor"] == "Liquidador Prueba"
    assert verify_chain(case.ledger.entries())[0]
    m = service.metrics(work)
    assert (m["tp"], m["tn"], m["fp"], m["fn"]) == (1, 1, 0, 0)
    with pytest.raises(ValueError):
        service.set_decision(case, "inventada", "x")


def test_web_import_networks_metrics_decision(tmp_path):
    import io as _io
    c = _web(tmp_path)
    csv_bytes = ("N° Siniestro;Fecha;RUT;Teléfono;Resultado\nW-1;01-05-2026 10:00;1-9;+56 9 1111 2222;fraude\n"
                 "W-2;02-06-2026 11:00;2-7;+56 9 1111 2222;pagado\nW-3;03-07-2026 12:00;3-5;+56 9 1111 2222;\n").encode("cp1252")
    r = c.post("/importar", data={"actor": "Prueba", "archivo": (_io.BytesIO(csv_bytes), "hist.csv")}, content_type="multipart/form-data")
    page = r.get_data(as_text=True)
    assert r.status_code == 200 and "3</b><span>siniestros importados" in page
    redes = c.get("/redes").get_data(as_text=True)
    assert "Red de 3 siniestros" in redes and "mismo teléfono" in redes
    met = c.get("/metricas").get_data(as_text=True)
    assert "Aciertos y errores" in met
    r = c.post("/caso/W-3/decision", data={"decision": "fraude", "actor": "Ana", "nota": "ok"}, follow_redirects=True)
    assert "Decisión registrada: Fraude confirmado" in r.get_data(as_text=True)
    assert "Fraude confirmado" in c.get("/siniestros").get_data(as_text=True)
    assert c.get("/plantilla.csv").get_data(as_text=True).lstrip("﻿").startswith("numero;fecha_siniestro")
    r = c.post("/importar", data={"archivo": (_io.BytesIO(b"x"), "hist.txt")}, content_type="multipart/form-data", follow_redirects=True)
    assert "Seleccione un archivo CSV o Excel" in r.get_data(as_text=True)


# ---- investigaciones: entrevistas, expediente y auditoría ------------------------------------
from veritas import interviews as IV, investigation as INV, report_audit as RA  # noqa: E402

INV_DEMO = ROOT / "demo" / "investigacion_demo"


def test_transcript_parsing_formats():
    numbered = "1. ¿Quién usaba habitualmente el vehículo?\nSolamente yo.\n2. ¿Algo más\nque agregar?\nNo, nada más.\n"
    items = IV.parse_text(numbered)
    assert [(q.n, q.a) for q in items] == [(1, "Solamente yo."), (2, "No, nada más.")] and items[1].q == "¿Algo más que agregar?"
    pr = IV.parse_text("P: ¿Qué hizo el otro vehículo?\nR: Se fugó del lugar.\nP: ¿Hora?\nR: A las 10:40.")
    assert len(pr) == 2 and pr[0].a == "Se fugó del lugar."


def test_interview_contradictions_and_documents():
    iv1 = IV.Interview("Ana", "Asegurada", items=IV.parse_text(
        "1. ¿Quién usaba habitualmente el vehículo?\nLo usamos los dos, yo y también mi hijo.\n"
        "2. ¿Qué pasó con el otro vehículo?\nSe fugó de inmediato.\n3. ¿Su RUT?\n11.111.111-1."), prefix="aaaaaaaa", index=1)
    iv2 = IV.Interview("Luis", "Conductor", items=IV.parse_text(
        "1. ¿Quién utilizaba habitualmente el vehículo?\nSolamente yo.\n2. ¿El otro vehículo se detuvo?\nNo tengo más información, no puedo aclarar.\n"
        "3. ¿A qué hora ocurrió el siniestro?\nA las 10 de la mañana.\n4. ¿Kilometraje?\nTenía 29.000 kilómetros."), prefix="bbbbbbbb", index=2)
    f = IV.compare([iv1, iv2], "2026-09-11T10:40:00")
    rules = {x.rule for x in f}
    assert {"contradiction_usuario_habitual", "contradiction_otro_vehiculo", "contradiction_hora"} <= rules
    assert any(x.evidence == ["aaaaaaaa:1001", "bbbbbbbb:2001"] for x in f)
    doc = IV.doc_facts_from_text("nota.pdf", "cccccccc:1", "Kilometraje: 46.592 km. RUT vendedor: 25.984.979-9")
    d = {x.rule for x in IV.cross_documents([iv1, iv2], [doc])}
    assert d == {"km_mismatch", "third_party_docs"}
    # sin contradicción: mismos hechos contados igual
    iv3 = IV.Interview("Eva", "Testigo", items=IV.parse_text("1. ¿Quién usaba habitualmente el vehículo?\nSolamente yo."), index=3)
    assert not [x for x in IV.compare([iv2, iv3]) if x.rule == "contradiction_usuario_habitual"]


def test_audit_finds_seeded_errors_in_fictional_report():
    rep = RA.read_pdf(INV_DEMO / "informe_para_auditar.pdf")
    assert len(rep.interviews) == 2 and len(rep.alerts) == 4 and rep.recommendation == "SE RECOMIENDA RECHAZAR EL RECLAMO"
    rules = {f.rule for f in RA.audit(rep, "deadbeef")}
    assert {"internal_trayecto", "contradiction_usuario_habitual", "contradiction_otro_vehiculo", "transcript_role",
            "weak_rejection", "empty_row", "name_variants"} <= rules
    assert not any("Página" in qa.a for iv in rep.interviews for qa in iv.items)


def test_expediente_flow_review_and_report(tmp_path):
    work = tmp_path / "w"
    done = INV.load_demo(work)
    assert done[0] == "SIN-DEMO-4471" and done[1].startswith("AUD-")
    case = Case(work / "SIN-DEMO-4471")
    a = INV.analyze(case)
    assert {"km_mismatch", "third_party_docs", "contradiction_usuario_habitual"} <= {f.rule for f in a["findings"]}
    assert len(a["verified"].accepted) >= 5 and not a["verified"].rejected
    review = {r.rule for r in INV.review(a)}
    assert {"review_missing", "review_unaddressed", "review_weak_rejection"} <= review
    # completar el trabajo: responder todo citando las contradicciones y concluir con un hecho acreditado
    all_ev = [e for f in a["findings"] for e in f.evidence]
    for i in range(1, 5):
        INV.save_response(case, i, "Acreditado" if i == 4 else "Confirmado", "Hallazgo", all_ev, "Analista")
    INV.save_conclusion(case, "Rechazar el reclamo", "Kilometraje adulterado acreditado con la nota de venta.", "Analista")
    a = INV.analyze(case)
    assert INV.review(a) == []
    html = INV.build_report(case, a, "Consultora Prueba")
    assert "CONSULTORA PRUEBA" in html and "RESPUESTA A LAS ALERTAS" in html.upper() and "15/09/2026 16:45" in html
    assert verify_chain(case.ledger.entries())[0]
    actions = [e["action"] for e in case.ledger.entries()]
    assert actions.count("alert_response") >= 6 and "conclusion" in actions


def test_web_investigations(tmp_path):
    import io as _io
    c = _web(tmp_path)
    assert "Aún no hay expedientes" in c.get("/investigaciones").get_data(as_text=True)
    c.post("/investigaciones/ejemplo")
    page = c.get("/investigacion/SIN-DEMO-4471").get_data(as_text=True)
    assert "Revisión previa" in page and "Qué dice cada declarante" in page and "E1·P9" in page
    r = c.post("/investigacion/SIN-DEMO-4471/alerta/2", data={"estado": "Acreditado", "hallazgo": "Bono confirma la atención",
                                                               "evidencia": ["x"], "actor": "Ana"}, follow_redirects=True)
    assert "Respuesta a la alerta 2 guardada" in r.get_data(as_text=True)
    assert "INFORME DE INVESTIGACIÓN" in c.get("/investigacion/SIN-DEMO-4471/informe").get_data(as_text=True)
    z = c.get("/investigacion/SIN-DEMO-4471/exportar")
    assert "verify.py" in zipfile.ZipFile(_io.BytesIO(z.data)).namelist()
    with open(INV_DEMO / "informe_para_auditar.pdf", "rb") as fh:
        r = c.post("/auditar", data={"actor": "Rev", "pdf": (fh, "informe.pdf")}, content_type="multipart/form-data", follow_redirects=True)
    assert "ya estaba auditado" in r.get_data(as_text=True) and "El informe se contradice" in r.get_data(as_text=True)
    r = c.post("/investigaciones/nuevo", data={"numero": "INV-1", "fecha_ocurrencia": "2026-09-01T10:00", "actor": "Ana",
                                              "alertas": "posible uso comercial", "iv_declarante": ["Uno", ""], "iv_rol": ["Asegurado", ""],
                                              "iv_fecha": ["", ""], "iv_texto": ["1. ¿Quién usaba habitualmente el vehículo?\nSolamente yo.", ""]},
               follow_redirects=True)
    assert "Expediente INV-1" in r.get_data(as_text=True)


def test_photo_metadata_tampering(tmp_path):
    """Fechas internas cambiadas, hora GPS que no calza y miniatura de otra imagen."""
    import io
    piexif = pytest.importorskip("piexif")
    from PIL import Image, ImageDraw
    from veritas.claims import quick_check

    def make(seed, dto, dt=None, thumb_from=None):
        img = Image.new("RGB", (800, 600), (90, 110, 130))
        d = ImageDraw.Draw(img)
        for i in range(20):
            x, y = (i * 37 * seed) % 700, (i * 53) % 500
            d.rectangle([x, y, x + 90, y + 70], fill=((i * 40) % 255, (i * 70 * seed) % 255, 60))
        img = Image.blend(img, Image.effect_noise((800, 600), 30).convert("RGB"), 0.1)  # textura de foto real
        t = (thumb_from or img).copy(); t.thumbnail((160, 120)); b = io.BytesIO(); t.save(b, "JPEG")
        ex = piexif.dump({
            "0th": {piexif.ImageIFD.Make: b"samsung", piexif.ImageIFD.DateTime: (dt or dto).encode()},
            "Exif": {piexif.ExifIFD.DateTimeOriginal: dto.encode(), piexif.ExifIFD.DateTimeDigitized: (dt or dto).encode(),
                     piexif.ExifIFD.OffsetTimeOriginal: b"-03:00", piexif.ExifIFD.ExposureTime: (1, 120),
                     piexif.ExifIFD.FNumber: (18, 10), piexif.ExifIFD.ISOSpeedRatings: 100},
            "GPS": {piexif.GPSIFD.GPSDateStamp: b"2026:09:20", piexif.GPSIFD.GPSTimeStamp: ((15, 1), (30, 1), (0, 1))},
            "1st": {}, "thumbnail": b.getvalue()})
        return img, ex

    def rules(name, *a, **k):
        img, ex = make(*a, **k)
        img.save(tmp_path / name, exif=ex)
        return {f.rule for f in quick_check(tmp_path / name, name, "2026-09-20T12:00:00")[1]}

    assert rules("ok.jpg", 1, "2026:09:20 12:30:00") == set()
    assert {"exif_dates_mismatch", "exif_gps_time_mismatch"} <= rules("a.jpg", 1, "2026:09:25 18:10:00", dt="2026:09:20 12:30:00")
    assert "exif_gps_time_mismatch" in rules("b.jpg", 1, "2026:09:25 18:10:00")
    orig, _ = make(1, "2026:09:20 12:30:00")
    assert "thumbnail_mismatch" in rules("c.jpg", 3, "2026:09:20 12:30:00", thumb_from=orig)
    Image.new("RGB", (1080, 2400), (250, 250, 250)).save(tmp_path / "p.png")
    assert "screenshot" in {f.rule for f in quick_check(tmp_path / "p.png", "p.png")[1]}


def _textured(seed, w=960, h=720):
    """Imagen con textura de foto real (ruido suave y fino), reproducible."""
    import numpy as np
    from PIL import Image
    rng = np.random.RandomState(seed)
    coarse = Image.fromarray(rng.randint(0, 255, (h // 40, w // 40, 3), dtype=np.uint8)).resize(
        (w, h), Image.Resampling.BICUBIC)
    arr = np.asarray(coarse, dtype=np.float32) * 0.6 + rng.randint(60, 190, 3) * 0.4
    arr += rng.normal(0, 6, arr.shape)
    return Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))


def test_pasted_region_and_clone_detected(tmp_path):
    import io
    from PIL import Image
    from veritas.image_content import analyze_image

    a = _textured(1)
    buf = io.BytesIO(); a.save(buf, "JPEG", quality=70); buf.seek(0)
    first = Image.open(buf).convert("RGB")
    clean = tmp_path / "limpia.jpg"; first.save(clean, quality=92)
    pasted = first.copy(); pasted.paste(_textured(2).crop((0, 0, 300, 240)), (330, 260))
    forged = tmp_path / "pegada.jpg"; pasted.save(forged, quality=92)
    cloned = a.copy(); cloned.paste(a.crop((60, 60, 260, 260)), (600, 380))
    clone_path = tmp_path / "clonada.jpg"; cloned.save(clone_path, quality=92)

    c = analyze_image(clean)
    assert not c["ghost"].get("region") and not c["clone"].get("src")
    r = analyze_image(forged)["ghost"].get("region")
    assert r, "no detectó la zona pegada"
    x0, y0, x1, y1 = r["bbox"]
    assert x0 < 630 and x1 > 330 and y0 < 500 and y1 > 260          # la zona marcada cae sobre lo pegado
    assert analyze_image(clone_path)["clone"].get("src"), "no detectó el clonado"


def test_solar_elevation_santiago():
    from datetime import datetime
    from veritas.image_content import solar_elevation
    # 20-09-2026 en Santiago: mediodía solar cerca de las 16:45 UTC, noche a las 02:00 UTC
    assert solar_elevation(-33.45, -70.66, datetime(2026, 9, 20, 16, 45)) > 50
    assert solar_elevation(-33.45, -70.66, datetime(2026, 9, 21, 2, 0)) < -20


def test_secure_capture_flow(tmp_path):
    import io
    from veritas import capture
    from veritas.web import create_app

    service.load_demo(tmp_path)
    case = service.find_case(tmp_path, LIMPIO)
    decl = service.declaration(case)
    app = create_app(tmp_path)
    c = app.test_client()
    assert c.post(f"/caso/{case.root.name}/captura").status_code == 302
    data = capture.load(case)
    tok = data["token"]
    assert c.get(f"/c/{tok}").status_code == 200
    assert c.get("/c/" + "x" * 30).status_code == 404
    buf = io.BytesIO(); _textured(5).save(buf, "JPEG", quality=92)
    r = c.post(f"/c/{tok}/foto", content_type="multipart/form-data", data={
        "shot": "frontal", "lat": str(decl["lat"]), "lon": str(decl["lon"]), "acc": "10",
        "foto": (io.BytesIO(buf.getvalue()), "x.jpg")})
    assert r.status_code == 200 and r.get_json()["ok"]
    bad = c.post(f"/c/{tok}/foto", content_type="multipart/form-data",
                 data={"shot": "otra", "foto": (io.BytesIO(buf.getvalue()), "x.jpg")})
    assert bad.status_code == 400
    entry = [e for e in case.ledger.entries() if e["action"] == "evidence_added"][-1]
    assert entry["data"]["capture"]["server_time"] and entry["data"]["capture"]["lat"] == decl["lat"]
    _, _, photos, _, findings, _, _ = service.run_analysis(case)
    cap = [p for p in photos if p.meta.get("secure_capture")]
    assert cap and cap[0].meta["taken"] == entry["data"]["capture"]["server_time"]
    assert not [f for f in findings if f.rule == "no_metadata" and cap[0].ev in f.evidence]
    # en modo red, desde otro equipo solo se puede abrir el enlace de captura
    app.config["LAN"] = True
    assert c.get("/", environ_base={"REMOTE_ADDR": "192.168.1.20"}).status_code == 403
    assert c.get(f"/c/{tok}", environ_base={"REMOTE_ADDR": "192.168.1.20"}).status_code == 200


# ---- diferenciadores forenses -------------------------------------------------------------
def test_chilean_validation():
    from veritas.chile import plate_format, plates_in, rut_dv, rut_valid, ruts_in
    assert rut_dv(12345678) == "5" and rut_valid("12.345.678-5") and rut_valid("12.345.678-9") is False
    assert rut_valid("sin rut") is None
    assert ruts_in("Taller RUT 76.543.210-3, cliente 12345678-9") == [("76.543.210-3", True), ("12.345.678-9", False)]
    assert plate_format("KXTR-45") == "nueva" and plate_format("AB1234") == "antigua" and plate_format("AEIO12") is None
    assert plates_in("patente KXTR·45 y BB-1234") == ["KXTR45", "BB1234"]


def test_whatsapp_parsing_android_and_iphone(tmp_path):
    import zipfile as zf
    from veritas.whatsapp import read
    text = ("19/09/26, 21:14 - Ana: mañana chocamos y\nle decimos al seguro\n"
            "[20/09/2026, 6:41:12 p. m.] Pedro: Se eliminó este mensaje\n"
            "20/09/26, 19:05 - Pedro: <Multimedia omitido>\n")
    p = tmp_path / "chat.txt"; p.write_text(text, encoding="utf-8")
    chat = read(p)
    assert [m.sender for m in chat.messages] == ["Ana", "Pedro", "Pedro"]
    assert chat.messages[0].text.endswith("le decimos al seguro")          # mensaje de varias líneas
    assert chat.messages[1].ts.hour == 18 and chat.messages[1].deleted     # 6:41 p. m.
    assert chat.messages[2].media
    z = tmp_path / "chat.zip"
    with zf.ZipFile(z, "w") as f:
        f.writestr("_chat.txt", text)
    assert len(read(z).messages) == 3
    (tmp_path / "nota.txt").write_text("hola, esto no es un chat", encoding="utf-8")
    assert read(tmp_path / "nota.txt") is None


def test_pdf_hidden_version_recovered(demo_src):
    from veritas.pdf_versions import extract, version_bytes
    raw = (demo_src / "actual" / "documentos" / "presupuesto_taller.pdf").read_bytes()
    r = extract(raw)
    assert [v["readable"] for v in r["versions"]] == [True, True]
    ch = r["changes"][0]
    assert ch["amounts"] == [["$1.850.000"], ["$4.650.000"]] and ch["dates"] == [["14-09-2026"], ["21-09-2026"]]
    assert ch["producer"] == "iLovePDF"
    v1 = version_bytes(raw, 1)
    assert v1.endswith(b"%%EOF\n") and len(v1) < len(raw)
    clean = (demo_src / "limpio" / "documentos" / "presupuesto_automotriz_central.pdf").read_bytes()
    assert extract(clean)["changes"] == []


def test_forensic_timeline_in_report_and_web(loaded):
    work, _ = loaded
    html = (work / ACTUAL / "informe.html").read_text(encoding="utf-8")
    assert "Línea de tiempo forense" in html and "traslado imposible" in html and "chat·" in html
    assert "qué cambió de la versión 1" in html and "$1.850.000" in html
    from veritas.web import create_app
    c = create_app(work).test_client()
    case = service.find_case(work, ACTUAL)
    digest = next(e["subject"] for e in case.ledger.entries()
                  if e["action"] == "evidence_added" and e["data"]["original_name"] == "presupuesto_taller.pdf")
    r = c.get(f"/caso/{case.root.name}/documento/{digest}/version/1")
    assert r.status_code == 200 and r.data.startswith(b"%PDF") and r.data.count(b"%%EOF") == 1
    assert c.get(f"/caso/{case.root.name}/documento/{digest}/version/9").status_code == 404
    assert c.get(f"/caso/{case.root.name}/documento/{'0' * 64}/version/1").status_code == 404


def test_external_services_with_mocked_provider(tmp_path, monkeypatch):
    from veritas import external
    from veritas.claims import Photo
    img = tmp_path / "evidence"; img.mkdir()
    (img / ("a" * 64)).write_bytes(b"\xff\xd8 imagen")

    class FakeCase:
        root, evidence_dir = tmp_path, img
    calls = []

    def fake_post(url, data, headers):
        calls.append(url)
        if "sightengine" in url:
            return {"status": "success", "type": {"ai_generated": 0.93, "ai_generators": {"midjourney": 0.9}}}
        return {"responses": [{"webDetection": {"fullMatchingImages": [{"url": "https://ejemplo.cl/auto.jpg"}],
                                                "pagesWithMatchingImages": [{"url": "https://ejemplo.cl/venta", "pageTitle": "Venta"}]}}]}
    monkeypatch.setattr(external, "_post", fake_post)
    photo = Photo("a" * 64, "foto.jpg", {}, "2026-09-20T18:00:00")
    cfg = {"sightengine_user": "u", "sightengine_secret": "s", "google_vision_key": "k"}
    res = external.run(FakeCase(), [photo], cfg)
    external.run(FakeCase(), [photo], cfg)                       # segunda vez: usa lo guardado, no vuelve a pagar
    assert len(calls) == 2
    rules = {f.rule: f for f in external.findings([photo], res)}
    assert rules["ai_detected"].severity == "alta" and "midjourney" in rules["ai_detected"].summary
    assert rules["found_online"].severity == "alta" and "ejemplo.cl/venta" in rules["found_online"].summary
    assert external.run(FakeCase(), [photo], {}) == {}


def test_chat_upload_accepted_only_if_real_chat(tmp_path):
    import io as _io
    c = _web(tmp_path)
    chat = "19/09/26, 21:14 - Ana: hola\n19/09/26, 21:15 - Pedro: chao\n".encode()
    r = c.post("/nuevo", data={"numero": "CH-1", "fecha_siniestro": "2026-09-20T18:30",
                                "fotos": [(_io.BytesIO(chat), "chat_whatsapp.txt")]}, content_type="multipart/form-data")
    assert r.status_code == 302
    r = c.post("/nuevo", data={"numero": "CH-2", "fecha_siniestro": "2026-09-20T18:30",
                                "fotos": [(_io.BytesIO(b"no es un chat"), "notas.txt")]}, content_type="multipart/form-data")
    assert r.status_code == 400 and "chats de WhatsApp" in r.get_data(as_text=True)
    assert "Servicios externos para fotos" in c.get("/configuracion").get_data(as_text=True)
    c.post("/configuracion", data={"empresa": "Consultora X", "google_vision_key": "k"})
    assert "Activos" in c.get("/configuracion").get_data(as_text=True)


# ---- lugares en chats, Excel, OCR y usuarios ----------------------------------------------
def test_places_in_chat_text():
    from veritas.places import mention
    assert mention("voy saliendo de Viña, llego tipo 8")[:1] == ("Viña del Mar",)
    assert mention("estoy en Rancagua con el auto")[0] == "Rancagua"
    assert mention("llegando a Temuco")[3] == "rumbo a"
    assert mention("Santiago me dijo que sí") is None and mention("el taller de Maipú cobra caro") is None


def test_chat_place_contradicts_crash(loaded):
    work, _ = loaded
    fs = [f for f in _analysis(work, ACTUAL)[4] if f.rule == "impossible_travel"]
    assert any("dice estar en Viña del Mar" in f.summary for f in fs)


def test_excel_exports(loaded):
    import io as _io
    from openpyxl import load_workbook
    from veritas.web import create_app
    c = create_app(loaded[0]).test_client()
    for what, sheet in (("cola", "Cola de trabajo"), ("redes", "Redes"), ("metricas", "Resumen")):
        r = c.get(f"/exportar/{what}.xlsx")
        assert r.status_code == 200 and sheet in load_workbook(_io.BytesIO(r.data)).sheetnames
    assert c.get("/exportar/otra.xlsx").status_code == 404


def test_ocr_reads_plate_and_scanned_document(loaded):
    from veritas import ocr
    if not ocr.available():
        pytest.skip("OCR no instalado (pip install rapidocr_onnxruntime)")
    work, _ = loaded
    _, _, photos, docs, findings, _, _ = _analysis(work, ACTUAL)
    rules = {f.rule: f for f in findings}
    assert "HJKL·21" in rules["plate_photo_mismatch"].summary and rules["plate_photo_mismatch"].severity == "alta"
    scanned = next(d for d in docs if d.name == "boleta_grua_escaneada.pdf")
    assert scanned.meta.get("ocr") and "18-09-2026" in scanned.meta["text"]
    assert any(f.rule == "doc_dated_before" and "boleta_grua" in f.summary for f in findings)
    assert not [f for f in _analysis(work, LIMPIO)[4] if f.rule == "plate_photo_mismatch"]


def test_users_login_roles_and_access_log(tmp_path):
    from veritas.web import create_app
    app = create_app(tmp_path / "casos")
    c = app.test_client()
    assert c.get("/").status_code == 200                                  # sin usuarios: modo demo abierto
    r = c.post("/usuarios", data={"usuario": "ana", "nombre": "Ana Pérez", "rol": "liquidador", "clave": "x" * 12},
               follow_redirects=True)
    assert "primer usuario debe ser administrador" in r.get_data(as_text=True)
    c.post("/usuarios", data={"usuario": "admin", "nombre": "Admin", "rol": "administrador", "clave": "clave-segura-1"})
    c.post("/usuarios", data={"usuario": "ana", "nombre": "Ana Pérez", "rol": "liquidador", "clave": "clave-segura-2"})
    c.post("/usuarios", data={"usuario": "ivan", "nombre": "Iván Inv", "rol": "investigador", "clave": "clave-segura-3"})
    users = json.loads((tmp_path / "casos" / "usuarios.json").read_text(encoding="utf-8"))
    assert "clave-segura-1" not in json.dumps(users) and set(users) == {"admin", "ana", "ivan"}
    c.get("/salir")
    r = c.get("/redes")
    assert r.status_code == 302 and "/ingresar" in r.location                # pide ingresar
    assert c.post("/ingresar", data={"usuario": "ana", "clave": "mala"}).status_code == 401
    r = c.post("/ingresar", data={"usuario": "ana", "clave": "clave-segura-2", "next": "/redes"})
    assert r.status_code == 302 and r.location.endswith("/redes")
    assert c.get("/redes").status_code == 200
    assert c.get("/metricas").status_code == 403 and c.get("/usuarios").status_code == 403   # rol liquidador
    c.post("/demo")
    case = service.find_case(tmp_path / "casos", LIMPIO)
    c.get(f"/caso/{case.root.name}")
    c.post(f"/caso/{case.root.name}/decision", data={"decision": "legitimo", "actor": "otro nombre"})
    assert service.current_decision(case)["actor"] == "Ana Pérez"           # el autor es el usuario, no lo escrito
    assert c.post("/demo", headers={"Origin": "https://sitio-malicioso.example"}).status_code == 403
    c.get("/salir")
    c.post("/ingresar", data={"usuario": "ivan", "clave": "clave-segura-3"})
    assert c.get("/investigaciones").status_code == 200 and c.get(f"/caso/{case.root.name}").status_code == 200
    assert c.post(f"/caso/{case.root.name}/decision", data={"decision": "fraude"}).status_code == 403
    c.get("/salir")
    for _ in range(5):
        c.post("/ingresar", data={"usuario": "admin", "clave": "mala"})
    r = c.post("/ingresar", data={"usuario": "admin", "clave": "clave-segura-1"})
    assert r.status_code == 401 and "Demasiados intentos" in r.get_data(as_text=True)
    log = (tmp_path / "casos" / "accesos.jsonl").read_text(encoding="utf-8")
    assert '"ver_caso"' in log and '"decision"' in log and '"ingreso_fallido"' in log


def test_heic_from_iphone_keeps_date(tmp_path):
    pytest.importorskip("pillow_heif")
    from PIL import Image
    from veritas.claims import photo_metadata
    exif = Image.Exif(); exif[0x010F] = "Apple"; exif[0x8769] = {0x9003: "2026:09:28 08:15:00"}
    p = tmp_path / "IMG_2001.HEIC"
    Image.new("RGB", (640, 480), (100, 110, 120)).save(p, format="HEIF", exif=exif.tobytes())
    m = photo_metadata(p, "IMG_2001.HEIC", content=False)
    assert m["make"] == "Apple" and m["taken"] == "2026-09-28T08:15:00"


def test_register_case_and_portal_for_insured(tmp_path):
    """El analista registra el caso sin evidencia, se genera el enlace y el asegurado sube fotos y documentos con título."""
    import io as _io
    from PIL import Image
    from veritas import capture
    c = _web(tmp_path)
    work = tmp_path / "casos"
    form = {"numero": "SIN-2026-2001", "fecha_siniestro": "2026-09-30T19:10", "asegurado": "Ana Pérez",
            "telefono": "+56 9 8123 4567", "email": "ana@example.com", "enviar_enlace": "1",
            "docs": ["licencia", "presupuesto"], "docs_sent": "1"}
    assert c.post("/nuevo", data=form, content_type="multipart/form-data").status_code == 302
    case = Case(work / "SIN-2026-2001")
    tok = capture.load(case)["token"]
    page = c.get("/caso/SIN-2026-2001").get_data(as_text=True)
    assert "https://wa.me/56981234567?text=" in page and "mailto:ana@example.com" in page
    assert "Esperando al asegurado" in c.get("/").get_data(as_text=True)
    portal = c.get(f"/c/{tok}").get_data(as_text=True)
    assert "Elegir o tomar fotos" in portal and "Elegir documentos" in portal and "Licencia de conducir" in portal

    def up(kind, name, blob, title):
        return c.post(f"/c/{tok}/archivo", data={"kind": kind, "titulo": title, "archivo": (_io.BytesIO(blob), name)},
                      content_type="multipart/form-data")
    pdf = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
    exif = Image.Exif(); exif[0x010F] = "Apple"; exif[0x0110] = "iPhone 14"; exif[0x8769] = {0x9003: "2026:09:30 19:14:00"}
    g = _io.BytesIO(); Image.new("RGB", (640, 480), (120, 110, 100)).save(g, "JPEG", exif=exif)
    raw = g.getvalue()
    r = up("photo", "IMG_1023.JPG", raw, "")
    assert r.status_code == 400 and "título" in r.get_json()["error"]          # el título es obligatorio
    assert not capture.progress(capture.load(case))["photos_ok"]
    r = up("photo", "IMG_1023.JPG", raw, "Parachoques trasero")
    assert r.status_code == 200 and r.get_json()["tipo"] == "photo" and r.get_json()["titulo"] == "Parachoques trasero"
    assert up("doc", "lic.pdf", pdf, "Licencia de conducir").get_json()["tipo"] == "doc"
    assert up("photo", "boleta.pdf", pdf, "Boleta grúa").get_json()["tipo"] == "doc"   # un PDF siempre es documento
    for kind, name, blob in (("photo", "virus.exe", b"MZ"), ("doc", "falso.pdf", b"hola"), ("otro", "x.pdf", pdf)):
        assert up(kind, name, blob, "Algo").status_code == 400
    pr = capture.progress(capture.load(case))
    assert pr["photos_ok"] and pr["docs_ok"] and (pr["done"], pr["total"]) == (2, 2)
    assert c.post(f"/c/{tok}/relato", data={"texto": "Iba por Apoquindo y me chocaron por atrás."}).status_code == 200
    assert c.post(f"/c/{tok}/terminar").status_code == 200
    assert c.post(f"/c/{tok}/terminar").status_code == 400                    # ya enviado: el enlace se cierra
    assert "Ya envió todo" in c.get(f"/c/{tok}").get_data(as_text=True)
    stored = next(e for e in case.ledger.entries() if e["action"] == "evidence_added"
                  and e["data"].get("portal", {}).get("title") == "Parachoques trasero")
    assert "parachoques_trasero" in stored["data"]["original_name"]
    assert (case.evidence_dir / stored["subject"]).read_bytes() == raw      # archivo original, byte a byte
    photo = next(p for p in service.run_analysis(case)[2] if "parachoques" in p.name)
    assert photo.meta["make"] == "Apple" and photo.meta["taken"] == "2026-09-30T19:14:00"
    case_page = c.get("/caso/SIN-2026-2001").get_data(as_text=True)
    assert "Entregado" in case_page and "Parachoques trasero" in case_page and "Licencia de conducir" in case_page
    assert "Parachoques trasero" in (work / "SIN-2026-2001" / "informe.html").read_text(encoding="utf-8")
    assert c.get("/c/" + "z" * 30).status_code == 404
    r = c.post("/nuevo", data={"numero": "X9", "fecha_siniestro": "2026-09-30T19:10"}, content_type="multipart/form-data")
    assert r.status_code == 400                                                 # sin archivos y sin enlace


def test_portal_paper_document_photo_is_read_as_document(tmp_path):
    """Una foto subida como documento (un papel) se lee como documento: RUT y fechas."""
    import io as _io
    from PIL import Image, ImageDraw, ImageFont
    from veritas import capture, ocr
    c = _web(tmp_path)
    form = {"numero": "SIN-2026-3001", "fecha_siniestro": "2026-09-30T19:10", "patente": "KXTR-45", "enviar_enlace": "1",
            "docs": ["presupuesto"], "docs_sent": "1", "docs_custom": "Carta del empleador; "}
    assert c.post("/nuevo", data=form, content_type="multipart/form-data").status_code == 302
    case = Case(tmp_path / "casos" / "SIN-2026-3001")
    tok = capture.load(case)["token"]
    assert "Carta del empleador" in c.get(f"/c/{tok}").get_data(as_text=True)
    img = Image.new("RGB", (1200, 1600), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.load_default(size=44)
    except TypeError:
        font = ImageFont.load_default()
    for i, line in enumerate(["TALLER EL ROBLE", "RUT: 76.543.210-9", "PRESUPUESTO", "Fecha: 25-09-2026", "TOTAL: $2.300.000"]):
        d.text((100, 150 + i * 90), line, fill="black", font=font)
    b = _io.BytesIO(); img.save(b, "JPEG", quality=90)
    r = c.post(f"/c/{tok}/archivo", data={"kind": "doc", "titulo": "Presupuesto en papel",
                                          "archivo": (_io.BytesIO(b.getvalue()), "IMG_3001.jpg")},
               content_type="multipart/form-data")
    assert r.status_code == 200
    _, _, photos, docs, findings, _, _ = service.run_analysis(case)
    assert photos == []                                        # la foto del papel no se trata como foto del vehículo
    paper = docs[0]
    assert paper.meta.get("image_doc") and paper.meta.get("title") == "Presupuesto en papel"
    if ocr.available():
        assert "76.543.210-9" in paper.meta["text"]
        assert {"doc_rut_invalid", "doc_dated_before"} <= {f.rule for f in findings}


def test_ocr_ignores_dates_that_look_like_plates():
    from veritas.ocr import plates
    assert plates([("DE 2026", 0.95), ("30 de septiembre de 2026", 0.9), ("AL1990", 0.9)]) == []
    assert plates([("KX·TR45", 0.95), ("HJ 1234", 0.9)]) == ["KXTR45", "HJ1234"]


def test_report_findings_show_thumbnail_and_guidance(loaded):
    work, _ = loaded
    html = (work / ACTUAL / "informe.html").read_text(encoding="utf-8")
    assert 'class="finding alta withpic"' in html and 'class="fthumb"' in html
    assert "Por qué importa:" in html and "Qué hacer:" in html


def test_public_link_warning_and_guard(tmp_path):
    from veritas import tunnel, capture
    assert tunnel.is_local("http://127.0.0.1:8765/c/x") and tunnel.is_local("https://192.168.1.5:8765/c/x")
    assert not tunnel.is_local("https://abc-def.trycloudflare.com/c/x")
    msg = capture.share_message("SIN-1", "Ana Pérez", "https://abc.trycloudflare.com/c/tok", "2026-10-08T10:00:00")
    assert "\nhttps://abc.trycloudflare.com/c/tok\n" in msg
    from veritas.web import create_app
    app = create_app(tmp_path / "casos")
    app.config.update(TESTING=True, PUBLIC=True, PUBLIC_BASE="https://abc.trycloudflare.com")
    cl = app.test_client()
    assert cl.get("/", headers={"Cf-Ray": "1"}).status_code == 404          # por el túnel solo se ve el portal
    assert cl.get("/siniestros", base_url="https://abc.trycloudflare.com").status_code == 404
    assert cl.get("/", base_url="http://127.0.0.1:8765").status_code in (200, 302)


def test_portal_uploads_show_in_case_and_report(tmp_path):
    import io, json
    from PIL import Image
    from veritas.web import create_app
    app = create_app(tmp_path / "casos"); app.config["TESTING"] = True
    cl = app.test_client(); B = "http://127.0.0.1:8765"; H = {"Origin": B}
    cl.post("/demo", base_url=B, headers=H)
    cid = [p.name for p in (tmp_path / "casos").iterdir() if (p / "ledger.jsonl").exists()][0]
    cl.post(f"/caso/{cid}/captura", data={"docs_sent": "1", "docs": ["presupuesto"]}, base_url=B, headers=H)
    cl.get(f"/caso/{cid}", base_url=B)
    root = tmp_path / "casos" / cid
    tok = json.loads((root / "captura.json").read_text())["token"]
    buf = io.BytesIO(); Image.new("RGB", (400, 300), (120, 90, 60)).save(buf, "JPEG")
    cl.post(f"/c/{tok}/archivo", data={"kind": "photo", "titulo": "parachoque", "archivo": (io.BytesIO(buf.getvalue()), "a.jpg")},
            content_type="multipart/form-data")
    h = cl.get(f"/caso/{cid}", base_url=B).get_data(as_text=True)      # sin que el asegurado presione Terminar
    assert "capthumbs" in h and "parachoque" in (root / "informe.html").read_text(encoding="utf-8")
    import re
    mini = re.search(r'/caso/[^"]+/archivo/[0-9a-f]{64}\?mini=1', h).group(0)
    assert cl.get(mini, base_url=B).mimetype == "image/jpeg"


def test_remove_wrong_file_keeps_custody(tmp_path):
    import io, json, re
    from PIL import Image
    from veritas.web import create_app
    from veritas.case import Case
    app = create_app(tmp_path / "casos"); app.config["TESTING"] = True
    cl = app.test_client(); B = "http://127.0.0.1:8765"; H = {"Origin": B}
    cl.post("/demo", base_url=B, headers=H)
    cid = [p.name for p in (tmp_path / "casos").iterdir() if (p / "ledger.jsonl").exists()][0]
    cl.post(f"/caso/{cid}/captura", data={"docs_sent": "1", "docs": ["presupuesto"]}, base_url=B, headers=H)
    root = tmp_path / "casos" / cid
    tok = json.loads((root / "captura.json").read_text())["token"]
    ids = []
    for t, c in (("foto equivocada", (10, 200, 10)), ("parachoque", (120, 90, 60))):
        buf = io.BytesIO(); Image.new("RGB", (400, 300), c).save(buf, "JPEG")
        r = cl.post(f"/c/{tok}/archivo", data={"kind": "photo", "titulo": t, "archivo": (io.BytesIO(buf.getvalue()), "a.jpg")},
                    content_type="multipart/form-data")
        ids.append(r.get_json()["id"])
    # el asegurado quita la equivocada
    assert cl.post(f"/c/{tok}/quitar", data={"id": ids[0]}).status_code == 200
    assert "foto equivocada" not in cl.get(f"/c/{tok}").get_data(as_text=True)
    case = Case(root)
    assert ids[0] in case.excluded() and (case.evidence_dir / ids[0]).exists()      # no se borra
    cl.get(f"/caso/{cid}", base_url=B)
    rep = (root / "informe.html").read_text(encoding="utf-8")
    assert "Archivo quitado del análisis" in rep and "parachoque" in rep
    assert 'class="photo"' in rep and "foto equivocada</b>" not in rep
    # el analista quita otra y luego la restaura
    cl.post(f"/caso/{cid}/archivo/{ids[1]}/quitar", data={"motivo": "Es de otro siniestro"}, base_url=B, headers=H)
    assert ids[1] in Case(root).excluded()
    assert "parachoque" not in cl.get(f"/c/{tok}").get_data(as_text=True)
    cl.post(f"/caso/{cid}/archivo/{ids[1]}/restaurar", base_url=B, headers=H)
    assert ids[1] not in Case(root).excluded()
    ok, _ = Case(root).verify_evidence(); assert ok
