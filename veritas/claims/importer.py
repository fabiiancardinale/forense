"""Importación masiva de siniestros históricos desde CSV o Excel.

Sirve para cargar el historial de la aseguradora y encontrar, de una vez, redes y patrones
que no se ven caso a caso. Cada fila se convierte en un siniestro con su propia cadena de
custodia, y el archivo original queda guardado con su hash para poder demostrar de dónde
salió cada dato.

Encabezados flexibles: se aceptan mayúsculas, tildes, espacios y nombres alternativos
("N° Siniestro", "Relato", "Monto", "Cuenta pago"...). Fechas en formato chileno
(20-09-2026, 20/09/2026 18:30) o ISO (2026-09-20). Separador coma o punto y coma.

Columna opcional `resultado` (fraude, pagado, rechazado, en revisión): se registra como la
decisión del liquidador y permite medir cuántos fraudes confirmados detecta Veritas.
"""
from __future__ import annotations

import csv
import io
import json
import re
import shutil
import tempfile
import unicodedata
from datetime import date, datetime
from pathlib import Path

from veritas.claims import service
from veritas.core.case import Case, sha256_file

TEMPLATE_COLUMNS = ["numero", "fecha_siniestro", "fecha_denuncia", "asegurado", "rut", "telefono", "email", "direccion",
                    "cuenta_bancaria", "patente", "lugar", "taller", "testigos", "parte_policial", "poliza",
                    "inicio_poliza", "fin_poliza", "cambio_cobertura", "suma_asegurada", "deducible",
                    "monto_reclamado", "descripcion", "resultado"]

ALIASES = {
    "numero": ["numero", "n_siniestro", "no_siniestro", "nro_siniestro", "numero_siniestro", "num_siniestro",
               "siniestro", "id_siniestro", "id"],
    "fecha_siniestro": ["fecha_siniestro", "fecha", "fecha_hora", "fecha_ocurrencia", "fecha_del_siniestro"],
    "hora": ["hora", "hora_siniestro"],
    "fecha_denuncia": ["fecha_denuncia", "fecha_aviso", "fecha_reporte", "denuncia"],
    "asegurado": ["asegurado", "nombre", "nombre_asegurado", "cliente"],
    "rut": ["rut", "rut_asegurado", "run"],
    "telefono": ["telefono", "fono", "celular", "movil"],
    "email": ["email", "correo", "mail", "correo_electronico"],
    "direccion": ["direccion", "domicilio", "direccion_asegurado"],
    "cuenta_bancaria": ["cuenta_bancaria", "cuenta", "cuenta_pago", "cuenta_de_pago", "n_cuenta"],
    "patente": ["patente", "placa", "ppu"],
    "lugar": ["lugar", "lugar_siniestro", "direccion_siniestro", "ubicacion"],
    "lat": ["lat", "latitud"], "lon": ["lon", "lng", "longitud"],
    "taller": ["taller", "taller_asignado", "servicio_tecnico"],
    "testigos": ["testigos", "n_testigos"],
    "parte_policial": ["parte_policial", "parte", "n_parte", "numero_parte"],
    "poliza": ["poliza", "n_poliza", "numero_poliza", "nro_poliza"],
    "inicio_poliza": ["inicio_poliza", "inicio_vigencia", "fecha_inicio", "vigencia_desde"],
    "fin_poliza": ["fin_poliza", "fin_vigencia", "vencimiento", "vigencia_hasta"],
    "cambio_cobertura": ["cambio_cobertura", "fecha_cambio_cobertura", "endoso", "fecha_endoso"],
    "suma_asegurada": ["suma_asegurada", "monto_asegurado", "capital_asegurado"],
    "deducible": ["deducible"],
    "monto_reclamado": ["monto_reclamado", "monto", "monto_siniestro", "monto_solicitado", "monto_reclamo"],
    "descripcion": ["descripcion", "relato", "detalle", "glosa", "descripcion_siniestro"],
    "resultado": ["resultado", "estado", "decision", "resolucion", "estado_final"],
}
DATE_FIELDS = ("fecha_denuncia", "inicio_poliza", "fin_poliza", "cambio_cobertura")


def _key(h: str) -> str:
    h = unicodedata.normalize("NFKD", str(h or "")).encode("ascii", "ignore").decode().lower()
    h = h.replace("n°", "n").replace("nº", "n")
    return re.sub(r"[^a-z0-9]+", "_", h).strip("_")


_LOOKUP = {_key(a): field for field, names in ALIASES.items() for a in names}


def map_headers(headers: list[str]) -> dict[str, str]:
    """Encabezado original -> campo de Veritas (solo los reconocidos)."""
    return {h: _LOOKUP[_key(h)] for h in headers if _key(h) in _LOOKUP}


def parse_datetime(value, with_time=True) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat() if with_time else value.date().isoformat()
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).isoformat() if with_time else value.isoformat()
    s = str(value).strip().replace("T", " ")
    s = re.sub(r"\s+", " ", s)
    formats = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
               "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d"]
    for fmt in formats:
        try:
            dt = datetime.strptime(s[:19], fmt)
            return dt.isoformat() if with_time else dt.date().isoformat()
        except ValueError:
            continue
    return None


def map_result(value) -> str | None:
    v = _key(value or "")
    if not v:
        return None
    if "fraud" in v:
        return "fraude"
    if any(x in v for x in ("legit", "pagad", "aprob", "indemniz", "cerrado_pago")):
        return "legitimo"
    if "rechaz" in v:
        return "rechazado"
    if "revision" in v or "pendiente" in v or "investig" in v:
        return "en_revision"
    return None


def read_table(path: Path) -> list[dict]:
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook  # pip install openpyxl
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [{h: v for h, v in zip(headers, r) if h} for r in rows[1:] if any(v not in (None, "") for v in r)]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    delim = ";" if sample.count(";") > sample.count(",") else ("\t" if sample.count("\t") > sample.count(",") else ",")
    return [r for r in csv.DictReader(io.StringIO(text), delimiter=delim) if any((v or "").strip() for v in r.values())]


def row_to_declaration(row: dict, mapping: dict[str, str]) -> tuple[dict | None, str | None, str | None]:
    """Devuelve (declaración, resultado, error)."""
    d = {}
    for original, field in mapping.items():
        v = row.get(original)
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        d[field] = v.strip() if isinstance(v, str) else v
    if not d.get("numero"):
        return None, None, "falta el número de siniestro"
    numero = str(d["numero"]).strip()
    if isinstance(d["numero"], float) and d["numero"].is_integer():
        numero = str(int(d["numero"]))
    d["numero"] = numero
    fecha = d.get("fecha_siniestro")
    if fecha is not None and d.get("hora") and not isinstance(fecha, datetime):
        fecha = f"{parse_datetime(fecha, with_time=False) or fecha} {d['hora']}"
    fecha = parse_datetime(fecha)
    if not fecha:
        return None, None, f"fecha de siniestro ausente o no reconocida ({d.get('fecha_siniestro', '')})"
    d["fecha_siniestro"] = fecha
    d.pop("hora", None)
    for f in DATE_FIELDS:
        if f in d:
            parsed = parse_datetime(d[f], with_time=False)
            if parsed:
                d[f] = parsed
            else:
                d.pop(f)
    for f in ("lat", "lon"):
        if f in d:
            try:
                d[f] = float(str(d[f]).replace(",", "."))
            except ValueError:
                d.pop(f)
    if ("lat" in d) != ("lon" in d):
        d.pop("lat", None), d.pop("lon", None)
    resultado = map_result(d.pop("resultado", None))
    return {k: (v if isinstance(v, (int, float)) else str(v)) for k, v in d.items()}, resultado, None


def import_history(workdir: Path, path: Path, actor: str = "importación", original_name: str | None = None) -> dict:
    """Crea un siniestro por fila y los analiza en orden cronológico. Devuelve un resumen."""
    workdir = Path(workdir)
    workdir.mkdir(parents=True, exist_ok=True)
    registry = workdir / "registro.jsonl"
    name = original_name or Path(path).name
    sha = sha256_file(path)
    store = workdir / "importaciones"
    store.mkdir(exist_ok=True)
    kept = store / f"{sha[:16]}{Path(name).suffix.lower()}"
    if not kept.exists():
        shutil.copy2(path, kept)
    rows = read_table(path)
    mapping = map_headers(list(rows[0].keys()) if rows else [])
    report = {"file": name, "sha256": sha, "rows": len(rows), "imported": [], "skipped": [],
              "recognized": sorted(set(mapping.values())), "ignored": sorted(h for h in (rows[0] if rows else {}) if h not in mapping)}
    if "numero" not in mapping.values() or "fecha_siniestro" not in mapping.values():
        report["error"] = "El archivo debe tener columnas de número de siniestro y fecha del siniestro."
        return report

    valid, seen = [], set()
    for i, row in enumerate(rows, start=2):  # fila 1 = encabezados
        decl, resultado, err = row_to_declaration(row, mapping)
        if err:
            report["skipped"].append((i, err))
            continue
        cid = service.case_id(decl["numero"])
        if cid in seen or (workdir / cid / "case.json").exists():
            report["skipped"].append((i, f"el siniestro {decl['numero']} ya existe"))
            continue
        seen.add(cid)
        valid.append((i, decl, resultado, cid))
    valid.sort(key=lambda x: x[1]["fecha_siniestro"])

    with tempfile.TemporaryDirectory() as td:
        for i, decl, resultado, cid in valid:
            dpath = Path(td) / "declaracion.json"
            dpath.write_text(json.dumps(decl, indent=2, ensure_ascii=False), encoding="utf-8")
            case = Case.create(workdir / cid, f"Siniestro {decl['numero']}", analyst=actor, kind="claim")
            case.add_evidence(dpath, analyst=actor, note="declaracion")
            case.ledger.append(actor, "imported_from", name, {"sha256": sha, "row": i})
            result = service.analyze_claim(case, registry)
            if resultado:
                service.set_decision(case, resultado, "historial", f"Resultado según {name}, fila {i}", source="importación")
            report["imported"].append((decl["numero"], result["level"], resultado))
    counts = {"bad": 0, "warn": 0, "ok": 0}
    for numero, _, _ in report["imported"]:  # estado final, después de los reanálisis por vínculo
        counts[(service.last_analysis(Case(workdir / service.case_id(numero))) or {}).get("level", "ok")] += 1
    report["levels"] = counts
    return report


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(TEMPLATE_COLUMNS)
    w.writerow(["SIN-2026-0001", "20-09-2026 18:30", "21-09-2026", "Nombre Apellido", "12.345.678-9", "+56 9 1234 5678",
                "correo@ejemplo.cl", "Calle 123, Comuna", "00-111-22222-33", "ABCD-12", "Av. Ejemplo 100, Comuna",
                "Taller Ejemplo", "1", "", "AUT-0001", "01-03-2025", "01-03-2027", "", "5.000.000", "150.000",
                "850.000", "Relato del asegurado", "pagado"])
    return "﻿" + buf.getvalue()  # BOM para que Excel respete las tildes
