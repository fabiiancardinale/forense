"""Exportación a Excel (.xlsx) de la cola de trabajo, las redes y las métricas."""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "0B1F36"
FILL = {"bad": "FEF3F2", "warn": "FFFAEB", "ok": "ECFDF3"}
TEXT = {"bad": "B42318", "warn": "B54708", "ok": "067647"}
MONEY = '"$"#,##0'


def _sheet(wb, title: str, headers: list[str], rows: list[list], widths: list[int], money_cols=(), level_col=None,
           note: str = ""):
    ws = wb.create_sheet(title)
    start = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="667588")
        start = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start, c, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start + 1):
        level = row[-1] if level_col is not None else None
        values = row[:-1] if level_col is not None else row
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            if c in money_cols and isinstance(v, (int, float)):
                cell.number_format = MONEY
            if level_col == c and level in FILL:
                cell.fill = PatternFill("solid", fgColor=FILL[level])
                cell.font = Font(bold=True, color=TEXT[level])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(start + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"
    return ws


def _bytes(wb) -> bytes:
    del wb[wb.sheetnames[0]]          # hoja vacía inicial
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stamp() -> str:
    return f"Generado por Veritas el {datetime.now():%d-%m-%Y %H:%M}"


def queue(cases: list[dict], level_text: dict) -> bytes:
    wb = Workbook()
    rows = [[c["numero"], c["patente"], c["asegurado"], c["fecha"], c["score"] if isinstance(c["score"], int) else None,
             c["findings"] if isinstance(c["findings"], int) else None, c["linked"] if isinstance(c["linked"], int) else None,
             c["monto"] or None, level_text.get(c["level"], c["rec"]), c["decision_label"] or "", c["level"]] for c in cases]
    _sheet(wb, "Cola de trabajo", ["N° siniestro", "Patente", "Asegurado", "Fecha", "Riesgo (0-100)", "Hallazgos", "Vinculados",
                                   "Monto reclamado", "Recomendación", "Decisión"], rows,
           [18, 11, 26, 17, 13, 11, 11, 16, 26, 24], money_cols=(8,), level_col=9,
           note=stamp() + ". El riesgo ordena la revisión; no es una probabilidad de fraude.")
    return _bytes(wb)


def networks(comps: list[dict], level_text: dict) -> bytes:
    wb = Workbook()
    summary = [[f"{'Red' if len(c['members']) >= 3 else 'Par'} {i}", len(c["members"]), c["people"], c["monto"] or None,
                (c["first"] or "")[:10], (c["last"] or "")[:10], ", ".join(f"{r} ×{n}" for r, n in c["reasons"])]
               for i, c in enumerate(comps, 1)]
    _sheet(wb, "Redes", ["Grupo", "Siniestros", "Asegurados", "Monto reclamado", "Desde", "Hasta", "Datos compartidos"],
           summary, [10, 11, 11, 16, 12, 12, 60], money_cols=(4,), note=stamp())
    detail = []
    for i, c in enumerate(comps, 1):
        for m in c["members"]:
            inf = c["info"][m]
            detail.append([f"{'Red' if len(c['members']) >= 3 else 'Par'} {i}", m, inf.get("asegurado", ""),
                           (inf.get("fecha") or "")[:10], inf.get("monto") or None, level_text.get(inf.get("level"), ""),
                           inf.get("decision_label") or "", inf.get("level")])
    _sheet(wb, "Siniestros en redes", ["Grupo", "N° siniestro", "Asegurado", "Fecha", "Monto", "Recomendación", "Decisión"],
           detail, [10, 18, 26, 12, 16, 26, 24], money_cols=(5,), level_col=6)
    return _bytes(wb)


def metrics(m: dict, rule_names: dict) -> bytes:
    wb = Workbook()
    rows = [["Fraudes confirmados detectados", f"{m['recall']}%", f"{m['tp']} de {m['tp'] + m['fn']}"],
            ["Alertas que resultaron fraude (precisión)", f"{m['precision']}%", f"{m['tp']} de {m['tp'] + m['fp']}"],
            ["Legítimos con alerta (falsas alarmas)", f"{m['false_alarm']}%", f"{m['fp']} de {m['fp'] + m['tn']}"],
            ["Monto en fraudes detectados", m["monto_detectado"], f"de ${m['monto_fraude']:,} en fraudes confirmados".replace(",", ".")],
            ["Siniestros en total", m["total"], ""]]
    ws = _sheet(wb, "Resumen", ["Indicador", "Valor", "Detalle"], rows, [42, 16, 40], note=stamp())
    ws.cell(7, 2).number_format = MONEY
    _sheet(wb, "Matriz", ["", "Veritas alertó", "Veritas no alertó"],
           [["Fraude confirmado", m["tp"], m["fn"]], ["Legítimo", m["fp"], m["tn"]]], [22, 16, 18])
    _sheet(wb, "Señales", ["Señal", "Siniestros"], [[rule_names.get(r, r), n] for r, n in m["rules"]], [44, 12])
    _sheet(wb, "Pendientes", ["N° siniestro", "Riesgo", "Monto"],
           [[r["numero"], r["score"], r["monto"] or None] for r in m["pendientes_alerta"]], [18, 10, 16], money_cols=(3,))
    return _bytes(wb)
